from pathlib import Path
from openpyxl import load_workbook


# ============================================================
# PHASE 5.3 — FINAL EXCEL WORKBOOK VALIDATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = PROJECT_ROOT / "reports" / "Billing_Assurance_Control_Tower.xlsx"

EXPECTED_SHEETS = [
    "Executive Summary",
    "Exception Queue",
    "Discount Approvals",
    "SLA Monitoring",
    "Control Summary",
    "Audit Trail",
]

EXPECTED_COUNTS = {
    "Exception Queue": 706,
    "Discount Approvals": 500,
    "SLA Monitoring": 500,
    "Control Summary": 5,
    "Audit Trail": 10,
}


print("=" * 70)
print("PHASE 5.3 — FINAL EXCEL WORKBOOK VALIDATION")
print("=" * 70)
print()

# ------------------------------------------------------------
# File check
# ------------------------------------------------------------

if not WORKBOOK_PATH.exists():
    raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

print("FILE CHECK")
print(f"  ✓ Workbook exists")
print(f"  ✓ File: {WORKBOOK_PATH}")
print(f"  ✓ Size: {WORKBOOK_PATH.stat().st_size:,} bytes")
print()

# ------------------------------------------------------------
# Open workbook
# ------------------------------------------------------------

wb = load_workbook(WORKBOOK_PATH, data_only=False)

# ------------------------------------------------------------
# Sheet validation
# ------------------------------------------------------------

print("SHEET VALIDATION")

actual_sheets = wb.sheetnames

for sheet in EXPECTED_SHEETS:
    if sheet not in actual_sheets:
        raise ValueError(f"Missing sheet: {sheet}")

    print(f"  ✓ {sheet}")

if len(actual_sheets) != len(EXPECTED_SHEETS):
    raise ValueError(
        f"Unexpected sheet count. "
        f"Expected {len(EXPECTED_SHEETS)}, found {len(actual_sheets)}"
    )

print()

# ------------------------------------------------------------
# Row count validation
# ------------------------------------------------------------

print("DATA COUNT VALIDATION")

for sheet_name, expected_count in EXPECTED_COUNTS.items():

    ws = wb[sheet_name]

    # One row is the header
    actual_count = ws.max_row - 1

    if actual_count != expected_count:
        raise ValueError(
            f"{sheet_name}: expected {expected_count}, "
            f"found {actual_count}"
        )

    print(
        f"  ✓ {sheet_name:<22} "
        f"{actual_count:,} rows"
    )

print()

# ------------------------------------------------------------
# Executive Summary validation
# ------------------------------------------------------------

print("EXECUTIVE SUMMARY VALIDATION")

ws = wb["Executive Summary"]

required_cells = [
    "B4",
    "B5",
    "B6",
    "B7",
    "B8",
    "B9",
    "B10",
    "B11",
]

for cell in required_cells:

    if ws[cell].value is None:
        raise ValueError(
            f"Executive Summary KPI cell is empty: {cell}"
        )

    print(
        f"  ✓ {cell}: {ws[cell].value}"
    )

print()

# ------------------------------------------------------------
# Formatting validation
# ------------------------------------------------------------

print("FORMATTING VALIDATION")

checks = {
    "Exception Queue": "A1",
    "Discount Approvals": "A1",
    "SLA Monitoring": "A1",
    "Control Summary": "A1",
    "Audit Trail": "A1",
}

for sheet_name, cell_ref in checks.items():

    ws = wb[sheet_name]
    cell = ws[cell_ref]

    if cell.fill.fill_type != "solid":
        raise ValueError(
            f"{sheet_name} header formatting missing"
        )

    if not cell.font.bold:
        raise ValueError(
            f"{sheet_name} header bold formatting missing"
        )

    print(
        f"  ✓ {sheet_name}: header formatting"
    )

print()

# ------------------------------------------------------------
# Freeze pane validation
# ------------------------------------------------------------

print("FREEZE PANE VALIDATION")

for sheet_name in EXPECTED_SHEETS[1:]:
    ws = wb[sheet_name]

    if ws.freeze_panes != "A2":
        raise ValueError(
            f"{sheet_name}: freeze pane expected A2, "
            f"found {ws.freeze_panes}"
        )

    print(
        f"  ✓ {sheet_name}: A2"
    )

print()

# ------------------------------------------------------------
# Final result
# ------------------------------------------------------------

print("=" * 70)
print("PHASE 5.3 COMPLETED SUCCESSFULLY")
print("=" * 70)
print()
print("EXCEL CONTROL WORKBOOK STATUS")
print("  ✓ Workbook exists")
print("  ✓ 6 required sheets")
print("  ✓ 706 exception records")
print("  ✓ 500 discount approval records")
print("  ✓ 500 SLA records")
print("  ✓ 5 control summary records")
print("  ✓ 10 audit records")
print("  ✓ Executive KPI section populated")
print("  ✓ Management formatting validated")
print("  ✓ Freeze panes validated")
print()
print("PHASE 5 — EXCEL CONTROL WORKBOOK COMPLETE")
print("=" * 70)