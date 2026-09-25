from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


# ============================================================
# PHASE 5.2 — ENHANCE EXCEL CONTROL WORKBOOK
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = PROJECT_ROOT / "reports" / "Billing_Assurance_Control_Tower.xlsx"


# ------------------------------------------------------------
# Load workbook
# ------------------------------------------------------------

wb = load_workbook(WORKBOOK_PATH)


# ------------------------------------------------------------
# Theme
# ------------------------------------------------------------

header_fill = PatternFill("solid", fgColor="1F4E78")
section_fill = PatternFill("solid", fgColor="D9EAF7")
kpi_fill = PatternFill("solid", fgColor="EAF3F8")
white_font = Font(color="FFFFFF", bold=True)
title_font = Font(size=18, bold=True)
section_font = Font(size=12, bold=True)
kpi_label_font = Font(size=10, bold=True)
kpi_value_font = Font(size=16, bold=True)

thin_border = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2"),
)


# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================

ws = wb["Executive Summary"]

ws.freeze_panes = "A4"
ws.sheet_view.showGridLines = False

# Title
ws["A1"] = "Billing Assurance & Revenue Protection Control Tower"
ws["A1"].font = title_font
ws["A1"].alignment = Alignment(vertical="center")

ws.merge_cells("A1:H1")
ws.row_dimensions[1].height = 30

# Subtitle
ws["A2"] = "Operational control, exception management, discount decisions and SLA monitoring"
ws["A2"].font = Font(italic=True, color="666666")

ws.merge_cells("A2:H2")

# KPI labels
kpi_rows = {
    4: "Control Hits",
    5: "Unique Exceptions",
    6: "Unique Customers",
    7: "Financial Impact",
    8: "Open Exceptions",
    9: "High Severity",
    10: "Medium Severity",
    11: "Low Severity",
}

for row, label in kpi_rows.items():
    ws[f"A{row}"].fill = kpi_fill
    ws[f"A{row}"].font = kpi_label_font
    ws[f"A{row}"].border = thin_border

    ws[f"B{row}"].fill = kpi_fill
    ws[f"B{row}"].font = kpi_value_font
    ws[f"B{row}"].border = thin_border

# KPI number formats
ws["B7"].number_format = '$#,##0.00'
for cell in ["B4", "B5", "B6", "B8", "B9", "B10", "B11"]:
    ws[cell].number_format = '#,##0'


# ------------------------------------------------------------
# Control Overview
# ------------------------------------------------------------

ws["A14"] = "CONTROL OVERVIEW"
ws["A14"].fill = header_fill
ws["A14"].font = white_font

ws.merge_cells("A14:H14")

for cell in ws[15]:
    if cell.value is not None:
        cell.fill = section_fill
        cell.font = section_font
        cell.border = thin_border


# ------------------------------------------------------------
# SLA Overview
# ------------------------------------------------------------

ws["A23"] = "SLA OVERVIEW"
ws["A23"].fill = header_fill
ws["A23"].font = white_font

ws.merge_cells("A23:H23")


# ------------------------------------------------------------
# Column widths
# ------------------------------------------------------------

widths = {
    "A": 30,
    "B": 20,
    "C": 20,
    "D": 20,
    "E": 20,
    "F": 20,
    "G": 20,
    "H": 20,
}

for col, width in widths.items():
    ws.column_dimensions[col].width = width


# ============================================================
# 2. EXCEPTION QUEUE
# ============================================================

ws = wb["Exception Queue"]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

ws.row_dimensions[1].height = 24

# Format columns by header name
headers = {cell.value: cell.column for cell in ws[1]}

for name in ["financial_impact", "exposure", "rate_residual"]:
    if name in headers:
        col = headers[name]
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).number_format = '$#,##0.00'

if "severity" in headers:
    col_letter = ws.cell(1, headers["severity"]).column_letter

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="HIGH"'],
            fill=PatternFill("solid", fgColor="F4CCCC")
        )
    )

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="MEDIUM"'],
            fill=PatternFill("solid", fgColor="FFF2CC")
        )
    )

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="LOW"'],
            fill=PatternFill("solid", fgColor="D9EAD3")
        )
    )


# ============================================================
# 3. DISCOUNT APPROVALS
# ============================================================

ws = wb["Discount Approvals"]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

ws.row_dimensions[1].height = 24

headers = {cell.value: cell.column for cell in ws[1]}

currency_columns = [
    "current_monthly_charge",
    "customer_value",
    "discount_cost",
    "estimated_retention_value",
    "net_economic_impact",
]

for name in currency_columns:
    if name in headers:
        col = headers[name]
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).number_format = '$#,##0.00'

for name in ["requested_discount_pct", "benefit_cost_ratio"]:
    if name in headers:
        col = headers[name]
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).number_format = '0.00'


# Decision highlighting
if "decision" in headers:
    col_letter = ws.cell(1, headers["decision"]).column_letter

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="DECLINE"'],
            fill=PatternFill("solid", fgColor="F4CCCC")
        )
    )

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="MANAGER_REVIEW"'],
            fill=PatternFill("solid", fgColor="FFF2CC")
        )
    )

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="AUTO_APPROVE"'],
            fill=PatternFill("solid", fgColor="D9EAD3")
        )
    )


# ============================================================
# 4. SLA MONITORING
# ============================================================

ws = wb["SLA Monitoring"]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

ws.row_dimensions[1].height = 24

headers = {cell.value: cell.column for cell in ws[1]}

for name in ["sla_target_hours", "turnaround_hours", "breach_hours"]:
    if name in headers:
        col = headers[name]
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).number_format = '0.00'


if "sla_status" in headers:
    col_letter = ws.cell(1, headers["sla_status"]).column_letter

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="Breached"'],
            fill=PatternFill("solid", fgColor="F4CCCC")
        )
    )

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="At Risk"'],
            fill=PatternFill("solid", fgColor="FFF2CC")
        )
    )

    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{col_letter}2="Met"'],
            fill=PatternFill("solid", fgColor="D9EAD3")
        )
    )


# ============================================================
# 5. CONTROL SUMMARY
# ============================================================

ws = wb["Control Summary"]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

headers = {cell.value: cell.column for cell in ws[1]}

for name in ["financial_impact"]:
    if name in headers:
        col = headers[name]
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).number_format = '$#,##0.00'


# ============================================================
# 6. AUDIT TRAIL
# ============================================================

ws = wb["Audit Trail"]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

headers = {cell.value: cell.column for cell in ws[1]}

if "financial_impact" in headers:
    col = headers["financial_impact"]
    for row in range(2, ws.max_row + 1):
        ws.cell(row, col).number_format = '$#,##0.00'


# ============================================================
# 7. GENERAL FORMATTING — ALL SHEETS
# ============================================================

for ws in wb.worksheets:

    ws.sheet_view.showGridLines = False

    # Default row height
    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = max(
            ws.row_dimensions[row].height or 15,
            18
        )

    # Alignment
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=False
                )

    # Autofit with reasonable maximum width
    for column_index, column_cells in enumerate(ws.iter_cols(), start=1):
        column_letter = get_column_letter(column_index)

        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            35
        )


# Re-apply wider Executive Summary columns
ws = wb["Executive Summary"]

for col, width in widths.items():
    ws.column_dimensions[col].width = width


# ============================================================
# SAVE
# ============================================================

wb.save(WORKBOOK_PATH)


print("=" * 70)
print("PHASE 5.2 — EXCEL MANAGEMENT PRESENTATION")
print("=" * 70)
print()
print("WORKBOOK ENHANCED")
print(f"  File: {WORKBOOK_PATH}")
print()
print("ENHANCEMENTS")
print("  ✓ Executive KPI formatting")
print("  ✓ Professional headers")
print("  ✓ Freeze panes")
print("  ✓ Filters")
print("  ✓ Currency formatting")
print("  ✓ Decision highlighting")
print("  ✓ Severity highlighting")
print("  ✓ SLA status highlighting")
print("  ✓ Column width optimization")
print("  ✓ Gridlines removed")
print("  ✓ Management-friendly layout")
print()
print("PHASE 5.2 COMPLETED SUCCESSFULLY")
print("=" * 70)