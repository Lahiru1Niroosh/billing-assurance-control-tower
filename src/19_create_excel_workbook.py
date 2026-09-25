import duckdb
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

# ============================================================
# PHASE 5.1 — EXCEL CONTROL WORKBOOK
# ============================================================

DB_PATH = Path("data/warehouse/billing_assurance.duckdb")
OUTPUT_DIR = Path("reports")
OUTPUT_FILE = OUTPUT_DIR / "Billing_Assurance_Control_Tower.xlsx"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

print("=" * 70)
print("PHASE 5.1 — EXCEL CONTROL WORKBOOK")
print("=" * 70)

# ------------------------------------------------------------
# CONNECT TO DUCKDB
# ------------------------------------------------------------

con = duckdb.connect(str(DB_PATH))

# ------------------------------------------------------------
# LOAD DATA
# ------------------------------------------------------------

executive = con.execute("""
    SELECT *
    FROM controls.v_executive_kpis
""").fetchdf()

exceptions = con.execute("""
    SELECT *
    FROM controls.v_exception_queue
""").fetchdf()

control_summary = con.execute("""
    SELECT *
    FROM controls.v_control_performance
""").fetchdf()

sla = con.execute("""
    SELECT *
    FROM staging.discount_decisions_sla
""").fetchdf()

audit = con.execute("""
    SELECT *
    FROM controls.v_control_audit_history
""").fetchdf()

discounts = con.execute("""
    SELECT
        request_id,
        customer_id,
        request_date,
        requested_discount_pct,
        reason,
        current_monthly_charge,
        customer_value,
        discount_cost,
        estimated_retention_value,
        net_economic_impact,
        benefit_cost_ratio,
        value_tier,
        decision_rule,
        decision,
        approval_level,
        decision_reason,
        sla_target_hours,
        turnaround_hours,
        sla_status,
        breach_hours
    FROM staging.discount_decisions_sla
""").fetchdf()

con.close()

# ------------------------------------------------------------
# CREATE WORKBOOK
# ------------------------------------------------------------

wb = Workbook()

# Remove default sheet
default_sheet = wb.active
wb.remove(default_sheet)

# ------------------------------------------------------------
# THEME
# ------------------------------------------------------------

header_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

section_fill = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

high_fill = PatternFill(
    "solid",
    fgColor="F4CCCC"
)

medium_fill = PatternFill(
    "solid",
    fgColor="FCE5CD"
)

low_fill = PatternFill(
    "solid",
    fgColor="D9EAD3"
)

title_font = Font(
    bold=True,
    size=16
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

thin_border = Border(
    bottom=Side(style="thin", color="D9E1F2")
)

# ------------------------------------------------------------
# HELPER FUNCTIONS
# ------------------------------------------------------------

def style_headers(ws):
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        cell.border = thin_border


def autofit(ws):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = cell.value

            if value is not None:
                max_length = max(
                    max_length,
                    len(str(value))
                )

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            40
        )


def add_table(ws, table_name):
    if ws.max_row < 2:
        return

    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    table = Table(
        displayName=table_name,
        ref=ref
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    table.tableStyleInfo = style
    ws.add_table(table)


def dataframe_to_sheet(ws, dataframe):
    # Headers
    for col_num, column in enumerate(dataframe.columns, start=1):
        ws.cell(
            row=1,
            column=col_num,
            value=column
        )

    # Data
    for row_num, row in enumerate(
        dataframe.itertuples(index=False),
        start=2
    ):
        for col_num, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            elif getattr(value, "tzinfo", None) is not None:
                value = value.replace(tzinfo=None)
            ws.cell(
                row=row_num,
                column=col_num,
                value=value
            )

    style_headers(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit(ws)


# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================

ws = wb.create_sheet("Executive Summary")

ws["A1"] = "Billing Assurance & Revenue Protection Control Tower"
ws["A1"].font = title_font

ws["A3"] = "EXECUTIVE KPI"
ws["A3"].fill = section_fill
ws["A3"].font = Font(bold=True)

kpi_row = executive.iloc[0]

kpis = [
    ("Control Hits", kpi_row["control_hits"]),
    ("Unique Exceptions", kpi_row["unique_exceptions"]),
    ("Unique Customers", kpi_row["unique_customers"]),
    ("Financial Impact", kpi_row["financial_impact"]),
    ("Open Exceptions", kpi_row["open_exceptions"]),
    ("High Severity", kpi_row["high_severity_exceptions"]),
    ("Medium Severity", kpi_row["medium_severity_exceptions"]),
    ("Low Severity", kpi_row["low_severity_exceptions"]),
    ("C001 Duplicate Records Involved", kpi_row["duplicate_records_involved"]),
    ("C001 Duplicate Customer Cases", kpi_row["duplicate_customer_cases"]),
]

for row_num, (label, value) in enumerate(kpis, start=4):
    ws.cell(row=row_num, column=1, value=label)
    ws.cell(row=row_num, column=2, value=float(value))

ws["B7"].number_format = '$#,##0.00'

ws["D3"] = "CONTROL OVERVIEW"
ws["D3"].fill = section_fill
ws["D3"].font = Font(bold=True)

for row_num, row in enumerate(
    control_summary.itertuples(index=False),
    start=4
):
    ws.cell(row=row_num, column=4, value=row.control_id)
    ws.cell(row=row_num, column=5, value=row.control_name)
    ws.cell(row=row_num, column=6, value=row.control_hits)
    ws.cell(row=row_num, column=7, value=row.financial_impact)

    ws.cell(row=row_num, column=7).number_format = '$#,##0.00'

for cell in ws[3][3:7]:
    cell.font = Font(bold=True)

ws.freeze_panes = "A4"

# ------------------------------------------------------------
# SLA SECTION
# ------------------------------------------------------------

ws["A15"] = "SLA OVERVIEW"
ws["A15"].fill = section_fill
ws["A15"].font = Font(bold=True)

sla_summary = con = None

# Calculate directly from dataframe
total_requests = len(sla)
sla_met = int((sla["sla_status"] == "Met").sum())
sla_at_risk = int((sla["sla_status"] == "At Risk").sum())
sla_breached = int((sla["sla_status"] == "Breached").sum())

sla_met_rate = (
    sla_met / total_requests * 100
    if total_requests
    else 0
)

sla_items = [
    ("Total Requests", total_requests),
    ("SLA Met", sla_met),
    ("SLA At Risk", sla_at_risk),
    ("SLA Breached", sla_breached),
    ("SLA Met Rate %", sla_met_rate),
    ("Non-Breach Rate %", (sla_met + sla_at_risk) / total_requests * 100 if total_requests else 0),
]

for row_num, (label, value) in enumerate(
    sla_items,
    start=16
):
    ws.cell(row=row_num, column=1, value=label)
    ws.cell(row=row_num, column=2, value=value)

for row in range(16, 16 + len(sla_items)):
    cell = ws.cell(row=row, column=2)
    if "Rate %" in ws.cell(row=row, column=1).value:
        cell.number_format = "0.00%"
        cell.value = ws.cell(row=row, column=2).value / 100

ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 18
ws.column_dimensions["D"].width = 15
ws.column_dimensions["E"].width = 28
ws.column_dimensions["F"].width = 15
ws.column_dimensions["G"].width = 18

# ============================================================
# 2. EXCEPTION QUEUE
# ============================================================

ws = wb.create_sheet("Exception Queue")

dataframe_to_sheet(ws, exceptions)

# Currency
if "financial_impact" in exceptions.columns:
    col = exceptions.columns.get_loc("financial_impact") + 1

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=col).number_format = '$#,##0.00'

# Conditional formatting for severity
severity_col = (
    exceptions.columns.get_loc("severity") + 1
    if "severity" in exceptions.columns
    else None
)

if severity_col:
    letter = get_column_letter(severity_col)

    ws.conditional_formatting.add(
        f"{letter}2:{letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{letter}2="HIGH"'],
            fill=high_fill
        )
    )

    ws.conditional_formatting.add(
        f"{letter}2:{letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{letter}2="MEDIUM"'],
            fill=medium_fill
        )
    )

    ws.conditional_formatting.add(
        f"{letter}2:{letter}{ws.max_row}",
        FormulaRule(
            formula=[f'{letter}2="LOW"'],
            fill=low_fill
        )
    )

add_table(ws, "ExceptionQueue")

# ============================================================
# 3. DISCOUNT APPROVALS
# ============================================================

ws = wb.create_sheet("Discount Approvals")

dataframe_to_sheet(ws, discounts)

currency_columns = [
    "current_monthly_charge",
    "customer_value",
    "discount_cost",
    "estimated_retention_value",
    "net_economic_impact",
]

for column in currency_columns:
    if column in discounts.columns:
        col = discounts.columns.get_loc(column) + 1

        for row in range(2, ws.max_row + 1):
            ws.cell(
                row=row,
                column=col
            ).number_format = '$#,##0.00'

if "requested_discount_pct" in discounts.columns:
    col = discounts.columns.get_loc("requested_discount_pct") + 1

    for row in range(2, ws.max_row + 1):
        ws.cell(
            row=row,
            column=col
        ).number_format = '0"%"'

if "benefit_cost_ratio" in discounts.columns:
    col = discounts.columns.get_loc("benefit_cost_ratio") + 1

    for row in range(2, ws.max_row + 1):
        ws.cell(
            row=row,
            column=col
        ).number_format = '0.00'

add_table(ws, "DiscountApprovalQueue")

# ============================================================
# 4. SLA MONITORING
# ============================================================

ws = wb.create_sheet("SLA Monitoring")

sla_columns = [
    "request_id",
    "customer_id",
    "approval_level",
    "sla_target_hours",
    "turnaround_hours",
    "sla_status",
    "breach_hours",
    "request_month",
]

sla_output = sla[sla_columns].copy()

dataframe_to_sheet(ws, sla_output)

add_table(ws, "SLAMonitoring")

status_col = sla_output.columns.get_loc("sla_status") + 1
status_letter = get_column_letter(status_col)

ws.conditional_formatting.add(
    f"{status_letter}2:{status_letter}{ws.max_row}",
    FormulaRule(
        formula=[f'{status_letter}2="Breached"'],
        fill=high_fill
    )
)

ws.conditional_formatting.add(
    f"{status_letter}2:{status_letter}{ws.max_row}",
    FormulaRule(
        formula=[f'{status_letter}2="At Risk"'],
        fill=medium_fill
    )
)

ws.conditional_formatting.add(
    f"{status_letter}2:{status_letter}{ws.max_row}",
    FormulaRule(
        formula=[f'{status_letter}2="Met"'],
        fill=low_fill
    )
)

# ============================================================
# 5. CONTROL SUMMARY
# ============================================================

ws = wb.create_sheet("Control Summary")

dataframe_to_sheet(ws, control_summary)

if "financial_impact" in control_summary.columns:
    col = control_summary.columns.get_loc("financial_impact") + 1

    for row in range(2, ws.max_row + 1):
        ws.cell(
            row=row,
            column=col
        ).number_format = '$#,##0.00'

add_table(ws, "ControlSummary")

# ============================================================
# 6. AUDIT TRAIL
# ============================================================

ws = wb.create_sheet("Audit Trail")

dataframe_to_sheet(ws, audit)

if "financial_impact" in audit.columns:
    col = audit.columns.get_loc("financial_impact") + 1

    for row in range(2, ws.max_row + 1):
        ws.cell(
            row=row,
            column=col
        ).number_format = '$#,##0.00'

add_table(ws, "AuditTrail")

# ============================================================
# FINAL WORKBOOK FORMATTING
# ============================================================

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

    # Header row height
    if ws.max_row >= 1:
        ws.row_dimensions[1].height = 24

    # General alignment
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="center"
            )

# ------------------------------------------------------------
# SAVE
# ------------------------------------------------------------

wb.save(OUTPUT_FILE)

print("\nWORKBOOK CREATED")
print(f"  File: {OUTPUT_FILE}")

print("\nSHEETS CREATED")

for ws in wb.worksheets:
    print(f"  ✓ {ws.title}")

print("\nDATA COUNTS")
print(f"  Exception Queue     : {len(exceptions):,}")
print(f"  Discount Approvals  : {len(discounts):,}")
print(f"  SLA Monitoring      : {len(sla_output):,}")
print(f"  Control Summary     : {len(control_summary):,}")
print(f"  Audit Trail         : {len(audit):,}")

print("\n" + "=" * 70)
print("PHASE 5.1 COMPLETED SUCCESSFULLY")
print("=" * 70)